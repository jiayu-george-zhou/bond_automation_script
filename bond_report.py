import pandas as pd
import openpyxl
from openpyxl.styles import Font
from openpyxl.styles import Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Border, Side

filename = "现券市场交易情况总结日报_20260723.xlsx"

df = pd.read_excel(
    filename,
    sheet_name="机构净买入债券成交金额统计表",
    header=2,  # row index 2 (the third row) is the real header
)

date_part = filename.split("_")[1]
date_part = date_part.split(".")[0]

year = date_part[:4]
month = date_part[4:6]
day = date_part[6:]

report_date = year + "/" + month + "/" + day
print(report_date)

df = df.rename(columns={"Unnamed: 0": "机构"})
df["机构"] = df["机构"].ffill()

cols_to_fix = [
    "国债-新债(Treasury Bond On-the-Run)",
    "国债-老债(Treasury Bond Off-the-Run)",
    "政策性金融债-新债(Policy Financial Bond On-the-Run)",
    "政策性金融债-老债(Policy Financial Bond Off-the-Run)",
    "地方政府债(Local Government Bond)",
]
for col in cols_to_fix:
    df[col] = df[col].replace("-", 0)
    df[col] = pd.to_numeric(df[col])

df = df[df["期限"].notna()]

df["国债(Treasury Bonds)"] = (
    df["国债-新债(Treasury Bond On-the-Run)"]
    + df["国债-老债(Treasury Bond Off-the-Run)"]
)

df["政策性金融债(Policy Financial Bonds)"] = (
    df["政策性金融债-新债(Policy Financial Bond On-the-Run)"]
    + df["政策性金融债-老债(Policy Financial Bond Off-the-Run)"]
)

df_final = df[
    [
        "机构",
        "期限",
        "国债(Treasury Bonds)",
        "政策性金融债(Policy Financial Bonds)",
        "地方政府债(Local Government Bond)",
    ]
]

maturity_map = {
    "1年及1年以下(Less then 1Y,including 1Y)": "1年及1年以下(Less than 1Y)",
    "1-3年(1~3Y)": "1-3年(1~3Y)",
    "3-5年(3~5Y)": "3-5年(3~5Y)",
    "5-7年(5~7Y)": "5-7年(5~7Y)",
    "7-10年(7~10Y)": "7-10年(7~10Y)",
    "10-15年(10~15Y)": "10年以上(10Y+)",
    "15-20年(15~20Y)": "10年以上(10Y+)",
    "20-30年(20~30Y)": "10年以上(10Y+)",
    "30年以上(More then 30Y)": "10年以上(10Y+)",
}

df_final = df_final[df_final["期限"] != "合计(Total)"]
df_final["期限分组(Maturity Group)"] = df_final["期限"].map(maturity_map)

maturity_order = [
    "1年及1年以下(Less than 1Y)",
    "1-3年(1~3Y)",
    "3-5年(3~5Y)",
    "5-7年(5~7Y)",
    "7-10年(7~10Y)",
    "10年以上(10Y+)",
]

df_final["期限分组(Maturity Group)"] = pd.Categorical(
    df_final["期限分组(Maturity Group)"], categories=maturity_order, ordered=True
)

bank_order = [
    "大型银行（Large banks）",
    "中小型银行（Small and medium-sized banks）",
    "证券公司（Securities companies）",
    "保险公司（The insurance company）",
    "基金公司及产品（Fund companies and products）",
    "货币市场基金（Money Market Fund）",
    "理财子公司及理财类产品（Wealth management Subsidiary and Products）",
    "其他(Others)",
]

df_final["机构"] = pd.Categorical(
    df_final["机构"], categories=bank_order
)

grouped = df_final.groupby(["机构", "期限分组(Maturity Group)"])[
    [
        "国债(Treasury Bonds)",
        "政策性金融债(Policy Financial Bonds)",
        "地方政府债(Local Government Bond)",
    ]
].sum()

treasury_table = grouped["国债(Treasury Bonds)"].unstack()
policy_table = grouped["政策性金融债(Policy Financial Bonds)"].unstack()
government_table = grouped["地方政府债(Local Government Bond)"].unstack()

treasury_table["合计(Total)"] = treasury_table.sum(axis=1)
policy_table["合计(Total)"] = policy_table.sum(axis=1)
government_table["合计(Total)"] = government_table.sum(axis=1)

combined = pd.concat(
    {
        "国债(Treasury Bonds)" : treasury_table,
        "政策性金融债(Policy Financial Bonds)" : policy_table,
        "地方政府债(Local Government Bond)" : government_table
    }
)

print(combined)

wb = openpyxl.Workbook()
ws = wb.active

maturity_headers = [
    "1年及1年以下(Less than 1Y)",
    "1-3年(1~3Y)",
    "3-5年(3~5Y)",
    "5-7年(5~7Y)",
    "7-10年(7~10Y)",
    "10年以上(10Y+)",
    "合计(Total)"
]

bold = Font(bold=True)
center = Alignment(horizontal="center", vertical="center")
wrap = Alignment(wrap_text=True, vertical="center")

title_and_section_align = Alignment(horizontal="center", vertical="center")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
name_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)

for i, header in enumerate(maturity_headers):
    ws.cell(row=2, column = 2+i, value=header).alignment = header_align

total_column = 2 + len(maturity_order)

def write_section(table, section_name, start_row):
    ws.cell(row=start_row, column=1, value=section_name).font = bold
    ws.cell(row=start_row, column=1).alignment = title_and_section_align
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=total_column)

    row_num = start_row + 1
    for institution in bank_order:
        ws.cell(row=row_num, column=1, value=institution).alignment = name_align
        ws.cell(row=row_num, column=1).border = thin_border

        for i, maturity in enumerate(maturity_order):
            value = table.loc[institution, maturity]
            cell = ws.cell(row=row_num, column = 2+i, value=value)
            cell.border = thin_border

        total = table.loc[institution, "合计(Total)"]
        total_cell = ws.cell(row=row_num, column=total_column, value=total)
        total_cell.border = thin_border

        row_num += 1

    data_range = f"B{start_row}:{openpyxl.utils.get_column_letter(total_column)}{row_num-1}"

    largest = table.abs().to_numpy().max()

    color_scale = ColorScaleRule(
        start_type="num", start_value=-largest, start_color="63BE7B",
        mid_type="num", mid_value=0, mid_color="FFFFFF",
        end_type="num", end_value=largest, end_color="F8696B"
    )
    ws.conditional_formatting.add(data_range, color_scale)

    return row_num

ws.cell(row=1, column=1, value="机构净买入债券成交金额统计表")
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_column)
ws.cell(row=2, column=1, value=report_date)

next_row = write_section(treasury_table, "国债", 3)
next_row = write_section(policy_table, "政策性金融债", next_row)
next_row = write_section(government_table, "地方政府债", next_row)

ws.column_dimensions["A"].width = 40
for col_letter in ["B","C","D","E","F","G","H"]:
    ws.column_dimensions[col_letter].width = 12

ws.cell(row=1, column=1).font = bold
ws.cell(row=1, column=1).alignment = center

wb.save(f"formatted_report_{date_part}.xlsx")